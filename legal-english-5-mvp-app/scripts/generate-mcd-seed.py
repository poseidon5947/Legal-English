#!/usr/bin/env python3
"""Transform MCD v1.3.81 → Delivery Model seed (LC-001 only)."""

from __future__ import annotations

import hashlib
import json
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SOURCE = Path("/home/user/Pictures/E-learning/Master_Content_Database_v1.3.81_Reconciled_2026-09-01.xlsx")
OUT = ROOT / "data" / "mcd-seed.json"
CANONICAL = {"Corporate Law", "Contracts", "Employment Law"}


def as_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, bool):
        return "true" if value else "false"
    return str(value).strip()


def as_bool(value) -> bool:
    if isinstance(value, bool):
        return value
    return as_text(value).lower() in {"true", "1", "yes", "x"}


def rows(wb, name: str) -> list[dict]:
    ws = wb[name]
    headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
    out = []
    for r in range(2, ws.max_row + 1):
        vals = [ws.cell(r, c).value for c in range(1, ws.max_column + 1)]
        if all(v is None or str(v).strip() == "" for v in vals):
            continue
        out.append(dict(zip(headers, vals)))
    return out


def content_hash(payload: dict) -> str:
    encoded = json.dumps(payload, ensure_ascii=False, sort_keys=True, separators=(",", ":"))
    return hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def main() -> None:
    wb = openpyxl.load_workbook(SOURCE, data_only=True)
    terms = {row["TermID"]: row for row in rows(wb, "Terms")}
    topics = {row["TopicName"]: row for row in rows(wb, "Topics")}
    categories = {row["CategoryID"]: row for row in rows(wb, "Categories")}
    items = [
        row
        for row in rows(wb, "CollectionItems")
        if row["LearningCollectionID"] == "LC-001" and row["Status"] == "Approved"
    ]
    use_items = rows(wb, "UseItWithItems")
    context_items = rows(wb, "InContextItems")
    quizzes = rows(wb, "QuizItems")
    variants = rows(wb, "TermJurisdictionVariants")

    use_by: dict[str, list] = defaultdict(list)
    for row in use_items:
        if row["Status"] != "Approved":
            continue
        use_by[row["TermID"]].append(
            {
                "id": as_text(row["UseItWithItemID"]),
                "expression": as_text(row["Expression"]),
                "displayOrder": int(row["DisplayOrder"] or 0),
            }
        )
    for lst in use_by.values():
        lst.sort(key=lambda item: item["displayOrder"])

    context_by: dict[str, dict] = {}
    grouped_ctx: dict[str, list] = defaultdict(list)
    for row in context_items:
        if row["Status"] != "Approved":
            continue
        grouped_ctx[row["TermID"]].append(row)
    for term_id, lst in grouped_ctx.items():
        chosen = sorted(lst, key=lambda row: int(row["DisplayOrder"] or 0))[0]
        context_by[term_id] = {
            "id": as_text(chosen["InContextItemID"]),
            "exampleText": as_text(chosen["ExampleText"]),
            "jurisdiction": as_text(chosen["Jurisdiction"]),
            "displayOrder": int(chosen["DisplayOrder"] or 0),
        }

    quiz_by = {}
    for row in quizzes:
        if row["Status"] != "Approved":
            continue
        options = []
        for key in ("OptionA", "OptionB", "OptionC", "OptionD"):
            value = as_text(row.get(key))
            if value:
                options.append(value)
        quiz_by[row["TermID"]] = {
            "id": as_text(row["QuizItemID"]),
            "question": as_text(row["Prompt"]),
            "options": options,
            "optionA": as_text(row["OptionA"]),
            "optionB": as_text(row["OptionB"]),
            "optionC": as_text(row["OptionC"]),
            "optionD": as_text(row.get("OptionD")),
            "correctOption": as_text(row["CorrectOption"]) or "A",
            "explanation": as_text(row["Explanation"]),
            "displayOrder": int(row["DisplayOrder"] or 1),
            "mcdStatus": as_text(row["Status"]),
        }

    variant_by: dict[str, dict] = defaultdict(dict)
    for row in variants:
        if row["Status"] != "Approved":
            continue
        payload = {
            "id": as_text(row["VariantID"]),
            "variantTerm": as_text(row["VariantTerm"]),
            "definition": as_text(row["Definition"]),
            "displayOrder": int(row["DisplayOrder"] or 0),
        }
        variant_by[row["TermID"]][as_text(row["Jurisdiction"])] = payload

    delivery = []
    category_counts = {"Corporate Law": 0, "Contracts": 0, "Employment Law": 0}
    for item in sorted(items, key=lambda row: int(row["Position"])):
        raw = terms[item["TermID"]]
        topic = topics[raw["Topic"]]
        category = categories[topic["CategoryID"]]["CategoryName"]
        if category not in CANONICAL:
            raise SystemExit(f"Non-canonical category for {item['TermID']}: {category}")
        category_counts[category] += 1
        us = as_bool(raw["JurisdictionUS"])
        uk = as_bool(raw["JurisdictionUK"])
        variants_for = variant_by.get(item["TermID"], {})
        quiz = quiz_by[item["TermID"]]
        use = use_by[item["TermID"]]
        context = context_by[item["TermID"]]
        hash_payload = {
            "id": item["TermID"],
            "term": as_text(raw["Term"]),
            "definition": as_text(raw["Definition"]),
            "spanishEquivalent": as_text(raw["SpanishEquivalent"]),
            "civilLawEquivalent": as_text(raw["CivilLawEquivalent"]),
            "spanishSpeakerAlert": as_text(raw["SpanishSpeakerAlert"]),
            "category": category,
            "useItWith": use,
            "inContext": context,
            "quiz": quiz,
            "usVariant": variants_for.get("US"),
            "ukVariant": variants_for.get("UK"),
        }
        delivery.append(
            {
                "id": item["TermID"],
                "term": as_text(raw["Term"]),
                "definition": as_text(raw["Definition"]),
                "spanishEquivalent": as_text(raw["SpanishEquivalent"]),
                "civilLawEquivalent": as_text(raw["CivilLawEquivalent"]),
                "spanishSpeakerAlert": as_text(raw["SpanishSpeakerAlert"]),
                "category": category,
                "topic": as_text(raw["Topic"]),
                "displayOrder": int(item["Position"]),
                "jurisdictionUS": us,
                "jurisdictionUK": uk,
                "jurisdiction": "US/UK" if us and uk else "UK" if uk else "US",
                "audioUsPath": as_text(raw["AudioUS"]),
                "audioUkPath": as_text(raw["AudioUK"]),
                "usVariant": variants_for.get("US"),
                "ukVariant": variants_for.get("UK"),
                "useItWith": use,
                "inContext": context,
                "quiz": quiz,
                "mcdStatus": as_text(raw["Status"]),
                "published": False,
                "archived": False,
                "sourceEditorialVersion": as_text(raw["EditorialVersion"]),
                "sourceLastReviewedAt": as_text(raw["LastReviewedAt"]),
                "sourceWorkbookVersion": "v1.3.81",
                "sourceContentHash": content_hash(hash_payload),
                "partOfSpeech": as_text(raw["PartOfSpeech"]),
                "comparativeLawNote": as_text(raw["ComparativeLawNote"]),
                "pronunciation": "",
            }
        )

    if category_counts != {"Corporate Law": 10, "Contracts": 10, "Employment Law": 10}:
        raise SystemExit(f"Expected 10/10/10, got {category_counts}")
    if len(delivery) != 30:
        raise SystemExit(f"Expected 30 LC-001 terms, got {len(delivery)}")

    payload = {
        "source": {
            "workbook": SOURCE.name,
            "version": "v1.3.81",
            "collection": "LC-001",
            "status": "OPERATIONAL MASTER CANDIDATE — NOT YET DESIGNATED CURRENT",
            "importedAt": "2026-09-03",
        },
        "categories": [
            {"id": "CAT-CORP", "name": "Corporate Law"},
            {"id": "CAT-CONT", "name": "Contracts"},
            {"id": "CAT-EMP", "name": "Employment Law"},
        ],
        "counts": {
            "terms": 30,
            "corporateLaw": 10,
            "contracts": 10,
            "employmentLaw": 10,
            "useItWith": sum(len(term["useItWith"]) for term in delivery),
            "inContext": sum(1 for term in delivery if term["inContext"]),
            "quizzes": sum(1 for term in delivery if term["quiz"]),
            "quizzesThreeOptions": sum(1 for term in delivery if term["quiz"] and len(term["quiz"]["options"]) == 3),
            "quizzesFourOptions": sum(1 for term in delivery if term["quiz"] and len(term["quiz"]["options"]) == 4),
            "variants": sum(1 for term in delivery if term["usVariant"]) + sum(1 for term in delivery if term["ukVariant"]),
            "published": 0,
            "audioUs": sum(1 for term in delivery if term["audioUsPath"]),
            "audioUk": sum(1 for term in delivery if term["audioUkPath"]),
        },
        "terms": delivery,
    }
    OUT.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(OUT, payload["counts"])


if __name__ == "__main__":
    main()
