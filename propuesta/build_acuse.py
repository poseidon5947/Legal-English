"""Generates the contractor's acuse for MCD_Delivery_Mapping_para_Carlos_v1.0 (§10)."""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "Acuse_Delivery_Mapping_v1.0_Carlos_Angola.docx"

CONTRACTOR = "Carlos Andrés Angola Berrio"
CONTRACTOR_ID = "C.C. 1.118.201.113 expedida en Villanueva, Casanare"
DATE = "3 de septiembre de 2026"

NAVY = RGBColor(0x12, 0x2B, 0x45)
MUTED = RGBColor(0x5B, 0x6B, 0x7B)


def para(doc, text, size=10.5, bold=False, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p


def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(headers):
        cell = t.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        run.bold = True
        run.font.size = Pt(9.5)
        run.font.color.rgb = NAVY
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = ""
            run = cells[i].paragraphs[0].add_run(value)
            run.font.size = Pt(9.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t


def build():
    doc = Document()
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(56)
        section.top_margin = section.bottom_margin = Pt(50)
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10.5)

    para(doc, "ACUSE DEL CONTRATISTA", size=16, bold=True, color=NAVY, space_after=2)
    para(doc, "MCD Delivery Mapping para Carlos · v1.0 · 2 de septiembre de 2026", size=11, color=MUTED, space_after=2)
    para(doc, "Baseline: MCD v1.3.81 (Reconciled 2026-09-01) · Selección LC-001 · 30 Terms", size=10, color=MUTED, space_after=12)

    table(
        doc,
        ["Campo", "Registro"],
        [
            ["Contratista", f"{CONTRACTOR} · {CONTRACTOR_ID}"],
            ["Documento recibido", "MCD_Delivery_Mapping_para_Carlos_v1.0_2026-09-02.docx y Master_Content_Database_v1.3.81_Reconciled_2026-09-01.xlsx"],
            ["Fecha del acuse", DATE],
            [
                "Acuse",
                "Recibí y leí el Delivery Mapping v1.0 completo. Acepto su límite de autoridad (§ Límite): no modifica las cinco fuentes, no convierte el MCD en Current Master "
                "y no autoriza publicación ni producción. Implemento el mapping lógico tal como está escrito; los nombres físicos se documentan sin alterarlo (C-03). "
                "Ningún registro se considera production-ready hasta cerrar import, audio, playback, publication-readiness, entitlement, security y rollback (§9, STOP RELEASE).",
            ],
        ],
    )

    para(doc, "Observaciones y estado de los controles pendientes (§ Cierre requerido)", size=12, bold=True, color=NAVY, space_after=6)
    table(
        doc,
        ["ID", "Control", "Respuesta del contratista", "Estado"],
        [
            [
                "C-01",
                "Convención de nombres de audio",
                "Propongo: {TermID}_US.{ext} y {TermID}_UK.{ext}, TermID en mayúsculas exactamente como en el MCD; ext ∈ mp3, m4a, wav, ogg. "
                "Ejemplos: CON-001_US.mp3, EMP-009_UK.mp3. Un archivo por Term y jurisdicción; repetir el nombre reemplaza el asset. "
                "La carga masiva asocia por nombre y rechaza (con motivo) todo nombre que no coincida con un TermID existente. "
                "Almacenamiento: bucket privado term-audio, ruta {TermID}/us.{ext} | {TermID}/uk.{ext}; el Learner reproduce por URL firmada de corta duración.",
                "Propuesto · requiere aprobación de la Owner antes del Hito B",
            ],
            [
                "C-02",
                "Quiz 3–4 opciones",
                "Acuso por escrito: OptionD es nullable; el quiz admite tres o cuatro opciones; CorrectOption debe apuntar a una opción poblada. "
                "La Asunción C de la propuesta queda corregida en la revisión del 3 de septiembre de 2026 (18 quizzes de 4 opciones y 12 de 3 en el baseline).",
                "Acusado",
            ],
            [
                "C-03",
                "Physical schema",
                "Documentado en supabase/migrations/001–004: terms (id text = TermID, category con CHECK de tres valores canónicos, display_order, mcd_status, "
                "source_editorial_version, source_last_reviewed_at, source_workbook_version, source_content_hash, audio_us_path, audio_uk_path, us_variant/uk_variant/use_it_with/in_context jsonb), "
                "quiz_items (option_d nullable, correct_option CHECK A–D), user_term_progress (state new/learning/mastered), subscriptions, import_runs, billing_events. "
                "El mapping lógico no se altera; la matriz campo a campo se entrega con el documento de arquitectura del Hito A.",
                "Documentado",
            ],
            [
                "C-04",
                "External import",
                "Implementado y probado en el alpha: carga inicial (30/30, 10/10/10, 0 Published), reimportación idempotente (mismo ID + mismo hash → 30 unchanged, sin tocar timestamps editoriales), "
                "rechazo de archivo inválido con reporte hoja/fila/ID/campo/valor/regla/severidad, y rollback que restaura el snapshot previo del lote (counts e IDs equivalentes). "
                "La evidencia formal se entrega en el TEST REQUIRED del Hito B, en ambiente de prueba.",
                "En curso · evidencia en Hito B",
            ],
            [
                "C-05",
                "Audio",
                "Flujo listo para recibir los 30 AudioUS y el AudioUK de EMP-009: carga individual desde el editor del Term, carga masiva por nombre, reemplazo, retiro y excepciones listadas. "
                "Retirar AudioUS (o AudioUK en EMP-009) devuelve el Term a borrador. Playback US/UK en el Learner. Pendiente: recibir los archivos producidos por la Owner.",
                "Pendiente de audios",
            ],
            [
                "C-06",
                "Publication readiness",
                "La compuerta de publicación exige quiz completo (≥3 opciones, puntero válido), AudioUS y AudioUK aplicable, y categoría canónica. Se vuelve a correr después de import y audio. "
                "El Vocabulary Dependency Gate se ejecuta con la Owner en la revisión del Hito B.",
                "Pendiente · después de audio",
            ],
            [
                "C-07",
                "Release gates",
                "Entitlement: máquina de estados active / past_due (gracia de 5 días) / payment_failed / cancelled (acceso hasta fin de periodo) / expired, con reconciliación idempotente y "
                "verificación de firma del webhook de Mercado Pago. Security: RLS por tabla y pruebas diferenciales Owner vs dos Learners. Rollback: ver C-04. "
                "La evidencia de los tres cierra en el Hito C antes de producción.",
                "Pendiente · Hito C",
            ],
        ],
    )

    para(doc, "Referencias de la propuesta afectadas", size=12, bold=True, color=NAVY, space_after=6)
    para(
        doc,
        "Propuesta MVP · revisión del 3 de septiembre de 2026: numerales 2, 3, 5 y 6; Hito B; supuestos A, C y E. La Asunción E incorpora la convención C-01 y la regla de publicación bloqueada sin AudioUS. "
        "Este acuse no modifica precio, alcance ni plazo de la propuesta.",
        space_after=14,
    )

    para(doc, "_____________________________________________", color=MUTED, space_after=0)
    para(doc, CONTRACTOR, bold=True, space_after=0)
    para(doc, CONTRACTOR_ID, color=MUTED, space_after=0)
    para(doc, f"Villanueva, Casanare · {DATE}", color=MUTED, align=WD_ALIGN_PARAGRAPH.LEFT)

    doc.save(OUT)
    print(OUT, OUT.stat().st_size)


if __name__ == "__main__":
    build()
